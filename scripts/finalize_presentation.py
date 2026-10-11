import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def update_presentation():
    pptx_path = "GapFill_AXION_Submission.pptx"
    bak_path = "GapFill_AXION_Submission.pptx.bak"

    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"Cannot find {pptx_path}")

    # Create backup
    if not os.path.exists(bak_path):
        shutil.copy2(pptx_path, bak_path)
        print(f"Backed up to {bak_path}")

    prs = Presentation(pptx_path)

    # ==========================================
    # SLIDE 1: Cover Title & Team Branding
    # ==========================================
    slide1 = prs.slides[0]

    # Clean any previously added custom title shapes if script re-runs
    for s in list(slide1.shapes):
        if s.name in ["GapFill_Title_Badge", "GapFill_Cover_Card"]:
            sp = s._element
            sp.getparent().remove(sp)

    # Add dark gold-accented project badge at the bottom of slide 1
    # Slide is 12191999 x 6858000 EMU
    card_w = 8592000
    card_h = 1180000
    card_l = int((prs.slide_width - card_w) / 2) # Perfectly centered
    card_t = 5380000

    badge_shape = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        card_l, card_t, card_w, card_h
    )
    badge_shape.name = "GapFill_Cover_Card"
    badge_shape.fill.solid()
    badge_shape.fill.fore_color.rgb = RGBColor(0x0B, 0x0F, 0x19) # Deep dark obsidian
    badge_shape.line.color.rgb = RGBColor(0xE5, 0xA9, 0x3C)       # Axion warm gold
    badge_shape.line.width = Pt(1.5)

    tf1 = badge_shape.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Pt(16)
    tf1.margin_right = Pt(16)
    tf1.margin_top = Pt(10)
    tf1.margin_bottom = Pt(10)

    # Line 1: Title
    p0 = tf1.paragraphs[0]
    p0.text = ""
    p0.alignment = PP_ALIGN.CENTER
    r0 = p0.add_run()
    r0.text = "GapFill"
    r0.font.name = "Aptos"
    r0.font.size = Pt(25)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p0.space_after = Pt(2)

    # Line 2: Subtitle
    p1 = tf1.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "AI-Powered Knowledge Gap Engine & Exam Preparation Companion"
    r1.font.name = "Aptos"
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8) # Slate-400
    p1.space_after = Pt(4)

    # Line 3: Team & College
    p2 = tf1.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2a = p2.add_run()
    r2a.text = "Team Non Challengers"
    r2a.font.name = "Aptos"
    r2a.font.size = Pt(13)
    r2a.font.bold = True
    r2a.font.color.rgb = RGBColor(0xFA, 0xCC, 0x15) # Warm Gold

    r2b = p2.add_run()
    r2b.text = "   •   "
    r2b.font.name = "Aptos"
    r2b.font.size = Pt(13)
    r2b.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    r2c = p2.add_run()
    r2c.text = "IIITDM Kancheepuram"
    r2c.font.name = "Aptos"
    r2c.font.size = Pt(13)
    r2c.font.bold = True
    r2c.font.color.rgb = RGBColor(0xF1, 0xF5, 0xF9)

    print("✅ Slide 1 (Cover) styled with GapFill title and team badge.")

    # ==========================================
    # SLIDE 2: Official Team Details Table
    # ==========================================
    slide2 = prs.slides[1]

    # Clean any previously added custom shapes if script re-runs
    for s in list(slide2.shapes):
        if s.name.startswith("Val_") or s.name in ["Team_Members_Card", "Team_Members_Left", "Team_Members_Right", "Team_Members_Divider"]:
            sp = s._element
            sp.getparent().remove(sp)

    # Exact field coordinates corresponding to template labels:
    # Row 1: Team Name (y=2267712, line y=2542032)
    # Row 2: Clg Name (y=2834640, line y=3108960)
    # Row 3: Leader Name (y=3401568, line y=3675888)
    # Row 4: State (y=3968496, line y=4242816)
    # Row 5: Leader Mail ID (y=4535424, line y=4809744)
    rows_data = [
        ("Val_TeamName", "Non Challengers", 2267712),
        ("Val_ClgName", "IIITDM Kancheepuram", 2834640),
        ("Val_LeaderName", "Pratyush Singh", 3401568),
        ("Val_State", "Tamil Nadu", 3968496),
        ("Val_LeaderEmail", "cs25i1064@iiitdm.ac.in", 4535424),
    ]

    val_left = 3200000
    val_width = 7800000
    val_height = 292608

    for name, value, top_y in rows_data:
        tb = slide2.shapes.add_textbox(val_left, top_y, val_width, val_height)
        tb.name = name
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = Pt(0)
        tf.margin_right = Pt(0)
        tf.margin_top = Pt(0)
        tf.margin_bottom = Pt(0)
        p = tf.paragraphs[0]
        p.text = value
        p.font.name = "Aptos"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x23, 0x2D, 0x3C) # Exact template navy color

    print("✅ Slide 2 table rows filled (Team, College, Leader, State, Email).")

    # ==========================================
    # SLIDE 2: Lower Card for Team Members & Links
    # ==========================================
    # Matching Rounded Rectangle 66: left=685800, width=10835640
    members_card_l = 685800
    members_card_t = 5080000
    members_card_w = 10835640
    members_card_h = 1320000

    card2 = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        members_card_l, members_card_t, members_card_w, members_card_h
    )
    card2.name = "Team_Members_Card"
    card2.fill.solid()
    card2.fill.fore_color.rgb = RGBColor(0xFA, 0xFA, 0xFA) # Matches Rounded Rectangle 66
    card2.line.color.rgb = RGBColor(0xD2, 0xD2, 0xD2)       # Matches Rounded Rectangle 66
    card2.line.width = Pt(1)

    # Vertical divider inside card
    div = slide2.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        5900000, 5200000, 10160, 1080000
    )
    div.name = "Team_Members_Divider"
    div.fill.solid()
    div.fill.fore_color.rgb = RGBColor(0xDC, 0xDC, 0xDC)
    div.line.fill.background()

    # Left Column: Team Members (Pooraw Kumar Dhurwey, Spandan Saha, Surabhi Sri Vyshnavi)
    tb_left = slide2.shapes.add_textbox(960120, 5160000, 4800000, 1150000)
    tb_left.name = "Team_Members_Left"
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Pt(0)
    tf_l.margin_right = Pt(0)
    tf_l.margin_top = Pt(0)
    tf_l.margin_bottom = Pt(0)

    # Heading
    p_lh = tf_l.paragraphs[0]
    p_lh.text = "TEAM MEMBERS (CONTRIBUTORS)"
    p_lh.font.name = "Aptos"
    p_lh.font.size = Pt(11)
    p_lh.font.bold = True
    p_lh.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue accent
    p_lh.space_after = Pt(4)

    # Members list
    members = [
        ("Member 2", "Pooraw Kumar Dhurwey"),
        ("Member 3", "Spandan Saha"),
        ("Member 4", "Surabhi Sri Vyshnavi"),
    ]

    for label, member_name in members:
        p_m = tf_l.add_paragraph()
        r_lbl = p_m.add_run()
        r_lbl.text = f"{label}: "
        r_lbl.font.name = "Aptos"
        r_lbl.font.size = Pt(12.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(0x23, 0x2D, 0x3C)

        r_nam = p_m.add_run()
        r_nam.text = member_name
        r_nam.font.name = "Aptos"
        r_nam.font.size = Pt(12.5)
        r_nam.font.bold = False
        r_nam.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p_m.space_after = Pt(3)

    # Right Column: Project & Verification Links
    tb_right = slide2.shapes.add_textbox(6100000, 5160000, 5200000, 1150000)
    tb_right.name = "Team_Members_Right"
    tf_r = tb_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Pt(0)
    tf_r.margin_right = Pt(0)
    tf_r.margin_top = Pt(0)
    tf_r.margin_bottom = Pt(0)

    # Heading
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "PROJECT DETAILS & LIVE DEMO"
    p_rh.font.name = "Aptos"
    p_rh.font.size = Pt(11)
    p_rh.font.bold = True
    p_rh.font.color.rgb = RGBColor(0x25, 0x63, 0xEB) # Blue accent
    p_rh.space_after = Pt(4)

    # Project Name
    p_proj = tf_r.add_paragraph()
    r_plbl = p_proj.add_run()
    r_plbl.text = "Project: "
    r_plbl.font.name = "Aptos"
    r_plbl.font.size = Pt(12)
    r_plbl.font.bold = True
    r_plbl.font.color.rgb = RGBColor(0x23, 0x2D, 0x3C)
    r_pnam = p_proj.add_run()
    r_pnam.text = "GapFill — Intelligent Knowledge Gap Engine"
    r_pnam.font.name = "Aptos"
    r_pnam.font.size = Pt(12)
    r_pnam.font.bold = False
    r_pnam.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    p_proj.space_after = Pt(3)

    # Live Website
    p_web = tf_r.add_paragraph()
    r_wlbl = p_web.add_run()
    r_wlbl.text = "Live App: "
    r_wlbl.font.name = "Aptos"
    r_wlbl.font.size = Pt(11.5)
    r_wlbl.font.bold = True
    r_wlbl.font.color.rgb = RGBColor(0x05, 0x96, 0x69) # Green accent
    r_wurl = p_web.add_run()
    r_wurl.text = "https://pratyush-singh123.github.io/Study-Planner/"
    r_wurl.font.name = "Aptos"
    r_wurl.font.size = Pt(11.5)
    r_wurl.font.bold = False
    r_wurl.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    p_web.space_after = Pt(3)

    # GitHub Repository
    p_gh = tf_r.add_paragraph()
    r_ghlbl = p_gh.add_run()
    r_ghlbl.text = "GitHub: "
    r_ghlbl.font.name = "Aptos"
    r_ghlbl.font.size = Pt(11.5)
    r_ghlbl.font.bold = True
    r_ghlbl.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    r_ghurl = p_gh.add_run()
    r_ghurl.text = "https://github.com/Pratyush-Singh123/Study-Planner"
    r_ghurl.font.name = "Aptos"
    r_ghurl.font.size = Pt(11.5)
    r_ghurl.font.bold = False
    r_ghurl.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    print("✅ Slide 2 team members container and live links added.")

    # Save finalized presentation
    prs.save(pptx_path)
    print(f"🎉 Successfully saved finalized presentation to {pptx_path}!")

if __name__ == "__main__":
    update_presentation()
